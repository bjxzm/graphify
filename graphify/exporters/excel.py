"""Excel workbook export for graphify knowledge graphs.

The first sheet mirrors the six-column data-labeling layout used by the
reference workbook supplied for this feature.  Additional sheets preserve the
graph structure and provenance so the tabular view remains auditable.
"""
from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Iterable

import networkx as nx

from graphify.analyze import _node_community_map


_MAX_EXCEL_ROWS = 1_048_576
_MAX_CELL_CHARS = 32_767
_FORMULA_PREFIXES = ("=", "+", "-", "@")
_CONFIDENCE_SCORES = {"EXTRACTED": 1.0, "INFERRED": 0.5, "AMBIGUOUS": 0.2}


def _excel_modules():
    try:
        from openpyxl import Workbook
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as exc:  # pragma: no cover - exercised without the extra
        raise ImportError(
            "Excel export requires openpyxl. Install it with: "
            "pip install 'graphifyy[excel]'"
        ) from exc
    return {
        "Workbook": Workbook,
        "ILLEGAL_CHARACTERS_RE": ILLEGAL_CHARACTERS_RE,
        "Alignment": Alignment,
        "Border": Border,
        "Font": Font,
        "PatternFill": PatternFill,
        "Side": Side,
        "Table": Table,
        "TableStyleInfo": TableStyleInfo,
    }


def _safe_text(value: Any, illegal_characters_re) -> str:
    """Return an Excel-safe string without turning untrusted text into a formula."""
    if value is None:
        return ""
    text = illegal_characters_re.sub("", str(value).replace("\r\n", "\n").replace("\r", "\n"))
    if text.lstrip().startswith(_FORMULA_PREFIXES):
        text = "'" + text
    if len(text) > _MAX_CELL_CHARS:
        text = text[: _MAX_CELL_CHARS - 1] + "…"
    return text


def _date_value(data: dict[str, Any]) -> date | datetime | str:
    for key in ("date", "published", "published_at", "created_at", "updated_at", "year"):
        raw = data.get(key)
        if raw in (None, ""):
            continue
        if isinstance(raw, (date, datetime)):
            return raw
        text = str(raw).strip()
        if len(text) == 4 and text.isdigit():
            return text
        try:
            parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
            return parsed.date() if parsed.time() == datetime.min.time() else parsed
        except ValueError:
            return text
    return ""


def _unique_join(values: Iterable[Any], *, separator: str = "，") -> str:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        text = str(value).strip() if value not in (None, "") else ""
        if text and text not in seen:
            seen.add(text)
            result.append(text)
    return separator.join(result)


def _node_content(node_id: str, data: dict[str, Any], G: nx.Graph) -> str:
    narrative = next(
        (
            str(data[key]).strip()
            for key in ("content", "description", "summary", "text", "docstring")
            if data.get(key) not in (None, "")
        ),
        "",
    )
    details: list[str] = []
    source = data.get("source_file") or data.get("url")
    location = data.get("source_location")
    if source:
        source_text = str(source)
        if location:
            source_text += f" ({location})"
        details.append(f"来源：{source_text}")

    relation_counts: defaultdict[str, int] = defaultdict(int)
    for _, _, edge in G.edges(node_id, data=True):
        relation_counts[str(edge.get("relation") or "relates_to")] += 1
    if relation_counts:
        relation_text = "，".join(
            f"{relation} {count}" for relation, count in sorted(relation_counts.items())
        )
        details.append(f"关系：{relation_text}")
    return "\n".join(part for part in (narrative, *details) if part)


def _edge_endpoints(u: Any, v: Any, data: dict[str, Any]) -> tuple[Any, Any]:
    return data.get("_src", u), data.get("_tgt", v)


def _style_sheet(
    sheet, modules, *, widths: dict[str, float], wrap_columns: set[str], body_row_height: float = 36
) -> None:
    Font = modules["Font"]
    PatternFill = modules["PatternFill"]
    Alignment = modules["Alignment"]
    Border = modules["Border"]
    Side = modules["Side"]

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="等线", size=11, bold=True, color="FFFFFF")
    body_font = Font(name="等线", size=11, color="1F1F1F")
    thin = Side(style="thin", color="D9E2F3")
    header_border = Border(bottom=Side(style="medium", color="17365D"))

    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 24
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = header_border
    for row_index, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        sheet.row_dimensions[row_index].height = body_row_height
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=cell.column_letter in wrap_columns,
            )
            cell.border = Border(bottom=thin)
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def _add_table(sheet, modules, name: str) -> None:
    if sheet.max_row < 2 or sheet.max_column < 1:
        return
    Table = modules["Table"]
    TableStyleInfo = modules["TableStyleInfo"]
    table = Table(displayName=name, ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def _require_row_capacity(kind: str, count: int) -> None:
    if count + 1 > _MAX_EXCEL_ROWS:
        raise ValueError(
            f"Excel export has {count:,} {kind}; the .xlsx row limit is "
            f"{_MAX_EXCEL_ROWS - 1:,} data rows per sheet."
        )


def to_excel(
    G: nx.Graph,
    communities: dict[int, list[str]],
    output_path: str,
    *,
    community_labels: dict[int, str] | None = None,
    cohesion: dict[int, float] | None = None,
) -> dict[str, int]:
    """Export ``G`` to a styled, source-traceable Excel workbook.

    Returns row counts for the main workbook tables.
    """
    modules = _excel_modules()
    Workbook = modules["Workbook"]
    illegal_re = modules["ILLEGAL_CHARACTERS_RE"]
    labels = {int(k): str(v) for k, v in (community_labels or {}).items()}
    cohesion = {int(k): float(v) for k, v in (cohesion or {}).items()}
    node_community = _node_community_map(communities)

    nodes = sorted(
        G.nodes(data=True),
        key=lambda item: (
            labels.get(node_community.get(item[0], -1), ""),
            str(item[1].get("label", item[0])).casefold(),
            str(item[0]),
        ),
    )
    edges = list(G.edges(data=True))
    _require_row_capacity("nodes", len(nodes))
    _require_row_capacity("relationships", len(edges))

    wb = Workbook()
    annotation = wb.active
    annotation.title = "数据标注"
    annotation.append(["序号", "时间", "主要标签", "次要标签", "关联标签", "内容"])
    for index, (node_id, data) in enumerate(nodes, start=1):
        cid = node_community.get(node_id)
        community_name = labels.get(cid, f"Community {cid}") if cid is not None else ""
        secondary = _unique_join(
            (community_name, data.get("type"), data.get("file_type"))
        )
        related = _unique_join(
            sorted(
                (G.nodes[n].get("label", n) for n in G.neighbors(node_id) if n in G),
                key=lambda value: str(value).casefold(),
            )
        )
        annotation.append(
            [
                index,
                _date_value(data),
                _safe_text(data.get("label", node_id), illegal_re),
                _safe_text(secondary, illegal_re),
                _safe_text(related, illegal_re),
                _safe_text(_node_content(str(node_id), data, G), illegal_re),
            ]
        )
    for cell in annotation["B"][1:]:
        if isinstance(cell.value, (date, datetime)):
            cell.number_format = "yyyy-mm-dd"
    _style_sheet(
        annotation,
        modules,
        widths={"A": 8, "B": 14, "C": 30, "D": 34, "E": 48, "F": 90},
        wrap_columns={"C", "D", "E", "F"},
        body_row_height=60,
    )
    _add_table(annotation, modules, "GraphifyAnnotations")

    node_sheet = wb.create_sheet("节点")
    node_sheet.append(
        ["节点ID", "标签", "类型", "文件类型", "社区ID", "社区名称", "度数", "来源文件", "来源位置", "说明"]
    )
    for node_id, data in nodes:
        cid = node_community.get(node_id)
        node_sheet.append(
            [
                _safe_text(node_id, illegal_re),
                _safe_text(data.get("label", node_id), illegal_re),
                _safe_text(data.get("type"), illegal_re),
                _safe_text(data.get("file_type"), illegal_re),
                cid if cid is not None else "",
                _safe_text(labels.get(cid, f"Community {cid}") if cid is not None else "", illegal_re),
                int(G.degree(node_id)),
                _safe_text(data.get("source_file") or data.get("url"), illegal_re),
                _safe_text(data.get("source_location"), illegal_re),
                _safe_text(_node_content(str(node_id), data, G), illegal_re),
            ]
        )
    _style_sheet(
        node_sheet,
        modules,
        widths={"A": 28, "B": 30, "C": 16, "D": 14, "E": 10, "F": 24, "G": 10, "H": 42, "I": 18, "J": 70},
        wrap_columns={"A", "B", "F", "H", "I", "J"},
    )
    _add_table(node_sheet, modules, "GraphifyNodes")

    edge_sheet = wb.create_sheet("关系")
    edge_sheet.append(
        ["源节点ID", "源标签", "关系", "目标节点ID", "目标标签", "证据类型", "置信分", "权重", "来源文件", "来源位置"]
    )
    for u, v, data in edges:
        source, target = _edge_endpoints(u, v, data)
        source_data = G.nodes[source] if source in G else {}
        target_data = G.nodes[target] if target in G else {}
        confidence = str(data.get("confidence") or "EXTRACTED")
        score = data.get("confidence_score", _CONFIDENCE_SCORES.get(confidence, 1.0))
        edge_sheet.append(
            [
                _safe_text(source, illegal_re),
                _safe_text(source_data.get("label", source), illegal_re),
                _safe_text(data.get("relation") or "relates_to", illegal_re),
                _safe_text(target, illegal_re),
                _safe_text(target_data.get("label", target), illegal_re),
                _safe_text(confidence, illegal_re),
                float(score) if isinstance(score, (int, float)) else _safe_text(score, illegal_re),
                float(data.get("weight", 1.0)) if isinstance(data.get("weight", 1.0), (int, float)) else _safe_text(data.get("weight"), illegal_re),
                _safe_text(data.get("source_file") or data.get("url"), illegal_re),
                _safe_text(data.get("source_location"), illegal_re),
            ]
        )
    for cell in edge_sheet["G"][1:]:
        cell.number_format = "0.00"
    for cell in edge_sheet["H"][1:]:
        cell.number_format = "0.00"
    _style_sheet(
        edge_sheet,
        modules,
        widths={"A": 28, "B": 28, "C": 20, "D": 28, "E": 28, "F": 16, "G": 12, "H": 10, "I": 42, "J": 18},
        wrap_columns={"A", "B", "C", "D", "E", "I", "J"},
    )
    _add_table(edge_sheet, modules, "GraphifyRelationships")

    community_sheet = wb.create_sheet("社区")
    community_sheet.append(["社区ID", "社区名称", "节点数", "凝聚度", "成员标签"])
    for cid, member_ids in sorted(communities.items()):
        existing = [node_id for node_id in member_ids if node_id in G]
        member_labels = _unique_join(
            sorted((G.nodes[n].get("label", n) for n in existing), key=lambda value: str(value).casefold())
        )
        community_sheet.append(
            [
                cid,
                _safe_text(labels.get(cid, f"Community {cid}"), illegal_re),
                len(existing),
                cohesion.get(cid, ""),
                _safe_text(member_labels, illegal_re),
            ]
        )
    for cell in community_sheet["D"][1:]:
        if isinstance(cell.value, (int, float)):
            cell.number_format = "0.00"
    _style_sheet(
        community_sheet,
        modules,
        widths={"A": 10, "B": 30, "C": 12, "D": 12, "E": 90},
        wrap_columns={"B", "E"},
    )
    _add_table(community_sheet, modules, "GraphifyCommunities")

    source_nodes: defaultdict[str, list[str]] = defaultdict(list)
    for node_id, data in nodes:
        source = data.get("source_file") or data.get("url")
        if source:
            source_nodes[str(source)].append(str(data.get("label", node_id)))
    source_edges: defaultdict[str, int] = defaultdict(int)
    for _, _, data in edges:
        source = data.get("source_file") or data.get("url")
        if source:
            source_edges[str(source)] += 1
    all_sources = sorted(set(source_nodes) | set(source_edges), key=str.casefold)
    _require_row_capacity("sources", len(all_sources))
    source_sheet = wb.create_sheet("来源")
    source_sheet.append(["来源", "节点数", "关系数", "主要标签"])
    for source in all_sources:
        source_sheet.append(
            [
                _safe_text(source, illegal_re),
                len(source_nodes[source]),
                source_edges[source],
                _safe_text(_unique_join(sorted(source_nodes[source], key=str.casefold)), illegal_re),
            ]
        )
    _style_sheet(
        source_sheet,
        modules,
        widths={"A": 70, "B": 12, "C": 12, "D": 90},
        wrap_columns={"A", "D"},
    )
    _add_table(source_sheet, modules, "GraphifySources")

    metadata_sheet = wb.create_sheet("元数据")
    metadata_sheet.append(["字段", "值"])
    metadata_rows = [
        ("生成时间", datetime.now(timezone.utc).replace(microsecond=0).isoformat()),
        ("节点数", G.number_of_nodes()),
        ("关系数", G.number_of_edges()),
        ("社区数", len(communities)),
        ("图类型", "有向" if G.is_directed() else "无向"),
        ("构建提交", G.graph.get("built_at_commit", "")),
        ("说明", "数据标注表采用：序号、时间、主要标签、次要标签、关联标签、内容。其余工作表保留图结构与来源追踪。"),
    ]
    for key, value in metadata_rows:
        metadata_sheet.append([_safe_text(key, illegal_re), _safe_text(value, illegal_re) if isinstance(value, str) else value])
    _style_sheet(
        metadata_sheet,
        modules,
        widths={"A": 20, "B": 100},
        wrap_columns={"B"},
    )
    _add_table(metadata_sheet, modules, "GraphifyMetadata")

    target = Path(output_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target)
    return {
        "annotations": len(nodes),
        "nodes": len(nodes),
        "relationships": len(edges),
        "communities": len(communities),
        "sources": len(all_sources),
    }


__all__ = ["to_excel"]
